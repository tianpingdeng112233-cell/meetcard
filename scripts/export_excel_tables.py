#!/usr/bin/env python3
"""
Export xty meetcard Excel `Back end` sheet into structured JSON for the meetcard PWA.

Layout assumptions (verified against Xiao Meet card(1).xlsx):
- Cols A-B: Male  (bodyweight, IPF GL coeff). Header row 2, data rows 3-1653
- Cols C-D: Female (bodyweight, IPF GL coeff). Same row range
- Cols G-J (rows 7-14): formula constants by sex/equipment/event combo
- Cols L-O (rows 8-24): plate jump table (min load, max load, 2a jump, 3a jump)
- Cols Q-S (rows 2-7): RPE feedback to next-attempt percentage table

Usage:
    python3 scripts/export_excel_tables.py /path/to/Xiao_Meet_card.xlsx > data/tables.json
"""
import json
import sys
from openpyxl import load_workbook


def main(xlsx_path: str) -> None:
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb["Back end"]

    # ---- Section 1: IPF GL coefficient lookup (cols A-D) ----
    ipf_gl_male: dict[str, float] = {}
    ipf_gl_female: dict[str, float] = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        bw_m, coeff_m, bw_f, coeff_f = row[0], row[1], row[2], row[3]
        if bw_m is not None and coeff_m is not None:
            ipf_gl_male[f"{float(bw_m):.1f}"] = round(float(coeff_m), 9)
        if bw_f is not None and coeff_f is not None:
            ipf_gl_female[f"{float(bw_f):.1f}"] = round(float(coeff_f), 9)

    # ---- Section 2: Formula constants (cols G-J, rows 7-14) ----
    formula_constants: dict[str, dict[str, float]] = {}
    constants_layout = [
        ("male/raw/powerlifting", 7),
        ("male/raw/bench-only", 8),
        ("male/equipped/powerlifting", 9),
        ("male/equipped/bench-only", 10),
        ("female/raw/powerlifting", 11),
        ("female/raw/bench-only", 12),
        ("female/equipped/powerlifting", 13),
        ("female/equipped/bench-only", 14),
    ]
    for label, row_idx in constants_layout:
        row = next(ws.iter_rows(min_row=row_idx, max_row=row_idx, values_only=True))
        formula_constants[label] = {
            "C1": float(row[6]) if row[6] is not None else 0.0,
            "C2": float(row[7]) if row[7] is not None else 0.0,
            "C3": float(row[8]) if row[8] is not None else 0.0,
            "C4": float(row[9]) if row[9] is not None else 0.0,
        }

    # ---- Section 3: Plate jump table (cols L-O, rows 8 onward) ----
    plate_jumps = []
    for row in ws.iter_rows(min_row=8, max_row=50, values_only=True):
        min_load, max_load, jump2a, jump3a = row[11], row[12], row[13], row[14]
        if min_load is None:
            break
        plate_jumps.append({
            "minLoad": float(min_load),
            "maxLoad": float(max_load),
            "jump2a": float(jump2a),
            "jump3a": float(jump3a),
        })

    # ---- Section 4: RPE feedback to next-attempt % (cols Q-S, rows 2 onward) ----
    rpe_feedback = []
    for row in ws.iter_rows(min_row=2, max_row=20, values_only=True):
        feedback, first_pct, second_pct = row[16], row[17], row[18]
        if not feedback or first_pct is None:
            break
        rpe_feedback.append({
            "feedback": str(feedback).strip(),
            "firstAttemptPct": float(first_pct),
            "secondAttemptPct": float(second_pct),
        })

    output = {
        "_meta": {
            "source": "Xiao Meet card.xlsx (Back end sheet)",
            "schema_version": 1,
            "exported_by": "export_excel_tables.py",
        },
        "ipfGlCoefficient": {
            "male": ipf_gl_male,
            "female": ipf_gl_female,
        },
        "formulaConstants": formula_constants,
        "plateJumps": plate_jumps,
        "rpeFeedback": rpe_feedback,
    }

    json.dump(output, sys.stdout, ensure_ascii=False, indent=2)
    sys.stdout.write("\n")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.stderr.write("Usage: export_excel_tables.py <path-to-xlsx>\n")
        sys.exit(1)
    main(sys.argv[1])
